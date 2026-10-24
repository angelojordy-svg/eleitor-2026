"""Leitura da planilha de boletins de urna (BU) por local de votação.

A planilha é gerada fora do projeto (TSE + cadastro de locais) e traz, para um
município, uma aba por recorte: locais com coordenadas, índice das urnas,
resumo por urna/cargo e os votos de cada votável em cada urna.

Só a leitura mora aqui; a agregação para o mapa fica em
``services.mapa_urnas``.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Sequence


class PlanilhaError(Exception):
    """Planilha ausente, num formato inesperado ou sem as abas necessárias."""


@dataclass(frozen=True)
class LocalVotacao:
    zona: str
    codigo: str
    nome: str
    endereco: str
    bairro: str
    cep: str
    latitude: float
    longitude: float
    secoes: int
    aptos: int
    comparecimento: int
    abstencoes: int


@dataclass(frozen=True)
class Urna:
    zona: str
    secao: str
    codigo: str
    local: str
    bairro: str
    agrega: str
    identificacao: str
    aptos: int
    comparecimento: int
    abstencoes: int
    validos_gov: int


@dataclass(frozen=True)
class ResumoUrna:
    zona: str
    secao: str
    codigo: str
    local: str
    agrega: str
    cargo: str
    aptos: int
    comparecimento: int
    abstencoes: int
    nominais: int
    legenda: int
    validos: int
    brancos: int
    nulos: int
    total: int


@dataclass(frozen=True)
class Voto:
    zona: str
    secao: str
    codigo: str
    local: str
    agrega: str
    cargo: str
    tipo: str
    numero: str
    votavel: str
    partido: str
    votos: int
    aptos: int
    comparecimento: int
    abstencoes: int


@dataclass
class PlanilhaBU:
    caminho: Path
    meta: dict[str, Any] = field(default_factory=dict)
    locais: list[LocalVotacao] = field(default_factory=list)
    urnas: list[Urna] = field(default_factory=list)
    resumos: list[ResumoUrna] = field(default_factory=list)
    votos: list[Voto] = field(default_factory=list)


def _slug(texto: Any) -> str:
    """Normaliza cabeçalhos: sem acento, minúsculo, espaços simples."""

    bruto = unicodedata.normalize("NFKD", str(texto or ""))
    sem_acento = "".join(c for c in bruto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", sem_acento.lower()).strip()


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _inteiro(valor: Any) -> int:
    if valor is None or valor == "":
        return 0
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, (int, float)):
        return int(round(valor))
    bruto = str(valor).strip().replace("\xa0", " ")
    if not bruto:
        return 0
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    try:
        return int(round(float(bruto)))
    except ValueError:
        return 0


def _decimal(valor: Any) -> float:
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    bruto = str(valor).strip().replace("\xa0", " ")
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    try:
        return float(bruto)
    except ValueError:
        return 0.0


def _colunas(linha: Sequence[Any]) -> dict[str, int]:
    vistos: dict[str, int] = {}
    for indice, valor in enumerate(linha):
        chave = _slug(valor)
        if chave and chave not in vistos:
            vistos[chave] = indice
    return vistos


def _indice(cols: dict[str, int], *nomes: str) -> int | None:
    for nome in nomes:
        alvo = _slug(nome)
        if alvo in cols:
            return cols[alvo]
    return None


def _valor(linha: Sequence[Any], posicao: int | None) -> Any:
    if posicao is None or posicao >= len(linha):
        return None
    return linha[posicao]


def _tabela(
    planilha: Any, obrigatorios: Sequence[str], limite_busca: int = 15
) -> tuple[dict[str, int], Iterator[Sequence[Any]]]:
    """Acha a linha de cabeçalho e devolve colunas + iterator das linhas seguintes."""

    linhas = planilha.iter_rows(values_only=True)
    for _ in range(limite_busca):
        linha = next(linhas, None)
        if linha is None:
            break
        cols = _colunas(linha)
        if all(_slug(nome) in cols for nome in obrigatorios):
            return cols, linhas
    raise PlanilhaError(
        f"Aba '{planilha.title}' sem cabeçalho reconhecível "
        f"(esperava: {', '.join(obrigatorios)})."
    )


def _abrir(caminho: Path) -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependência declarada
        raise PlanilhaError(
            "openpyxl não instalado. Rode: pip install openpyxl"
        ) from exc
    try:
        return load_workbook(caminho, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise PlanilhaError(f"Não consegui abrir a planilha: {caminho} ({exc})") from exc


def _aba(wb: Any, *nomes: str) -> Any:
    procuradas = {_slug(nome) for nome in nomes}
    for planilha in wb.worksheets:
        if _slug(planilha.title) in procuradas:
            return planilha
    raise PlanilhaError(
        "A planilha não tem a aba "
        + " / ".join(f"'{nome}'" for nome in nomes)
        + f". Abas encontradas: {', '.join(p.title for p in wb.worksheets)}."
    )


def _ler_meta(planilha: Any) -> dict[str, Any]:
    pares: dict[str, str] = {}
    for linha in planilha.iter_rows(max_row=45, values_only=True):
        if not linha:
            continue
        chave = _slug(linha[0])
        if not chave or len(linha) < 2:
            continue
        valor = _texto(linha[1])
        if valor:
            pares.setdefault(chave, valor)

    municipio = pares.get("municipio", "")
    codigo_tse = ""
    ibge = ""
    codigos = re.search(r"(\d+)\s*/\s*(\d+)", pares.get("codigo tse ibge", ""))
    if codigos:
        codigo_tse, ibge = codigos.group(1), codigos.group(2)
    eleicao = pares.get("eleicao", "")
    ano = re.search(r"(20\d{2})", eleicao)
    turno = re.search(r"(\d+)\s*[ºo°]?\s*turno", eleicao)
    uf = municipio.rsplit("/", 1)[-1].strip().upper() if "/" in municipio else ""

    return {
        "municipio": municipio,
        "uf": uf,
        "codigo_tse": codigo_tse,
        "ibge": ibge,
        "eleicao": eleicao,
        "ano": int(ano.group(1)) if ano else None,
        "turno": int(turno.group(1)) if turno else None,
        "fonte": pares.get("fonte", ""),
        "gerado_tse": pares.get("geracao do arquivo do tse", ""),
        "data_planilha": pares.get("data desta planilha", ""),
        "resumo_zonas": pares.get("zonas e secoes", ""),
        "resumo_locais": pares.get("locais de votacao", ""),
        "cargos": pares.get("cargos", ""),
    }


def _ler_locais(planilha: Any) -> list[LocalVotacao]:
    cols, linhas = _tabela(
        planilha,
        ("zona", "local de votação", "latitude", "longitude", "aptos"),
    )
    c_zona = _indice(cols, "zona")
    c_codigo = _indice(cols, "código", "cód. local", "cod local")
    c_nome = _indice(cols, "local de votação")
    c_endereco = _indice(cols, "endereço")
    c_bairro = _indice(cols, "bairro")
    c_cep = _indice(cols, "cep")
    c_lat = _indice(cols, "latitude")
    c_lon = _indice(cols, "longitude")
    c_secoes = _indice(cols, "seções", "secoes")
    c_aptos = _indice(cols, "aptos")
    c_comp = _indice(cols, "comparecimento")
    c_abst = _indice(cols, "abstenções", "abstencoes")

    locais: list[LocalVotacao] = []
    for linha in linhas:
        zona = _texto(_valor(linha, c_zona))
        if not zona.isdigit():
            continue
        nome = _texto(_valor(linha, c_nome))
        if not nome:
            continue
        locais.append(
            LocalVotacao(
                zona=zona,
                codigo=_texto(_valor(linha, c_codigo)),
                nome=nome,
                endereco=_texto(_valor(linha, c_endereco)),
                bairro=_texto(_valor(linha, c_bairro)),
                cep=_texto(_valor(linha, c_cep)),
                latitude=_decimal(_valor(linha, c_lat)),
                longitude=_decimal(_valor(linha, c_lon)),
                secoes=_inteiro(_valor(linha, c_secoes)),
                aptos=_inteiro(_valor(linha, c_aptos)),
                comparecimento=_inteiro(_valor(linha, c_comp)),
                abstencoes=_inteiro(_valor(linha, c_abst)),
            )
        )
    return locais


def _ler_urnas(planilha: Any) -> list[Urna]:
    cols, linhas = _tabela(
        planilha,
        ("zona", "seção", "cód. local", "aptos", "comparecimento"),
    )
    c_zona = _indice(cols, "zona")
    c_secao = _indice(cols, "seção", "secao")
    c_codigo = _indice(cols, "cód. local", "código", "cod local")
    c_local = _indice(cols, "local de votação")
    c_bairro = _indice(cols, "bairro")
    c_agrega = _indice(cols, "agrega")
    c_urna = _indice(cols, "urna")
    c_aptos = _indice(cols, "aptos")
    c_comp = _indice(cols, "comparecimento")
    c_abst = _indice(cols, "abstenções", "abstencoes")
    c_validos = _indice(cols, "válidos gov", "validos gov")

    urnas: list[Urna] = []
    for linha in linhas:
        zona = _texto(_valor(linha, c_zona))
        secao = _texto(_valor(linha, c_secao))
        if not zona.isdigit() or not secao.isdigit():
            continue
        urnas.append(
            Urna(
                zona=zona,
                secao=secao,
                codigo=_texto(_valor(linha, c_codigo)),
                local=_texto(_valor(linha, c_local)),
                bairro=_texto(_valor(linha, c_bairro)),
                agrega=_texto(_valor(linha, c_agrega)),
                identificacao=_texto(_valor(linha, c_urna)),
                aptos=_inteiro(_valor(linha, c_aptos)),
                comparecimento=_inteiro(_valor(linha, c_comp)),
                abstencoes=_inteiro(_valor(linha, c_abst)),
                validos_gov=_inteiro(_valor(linha, c_validos)),
            )
        )
    return urnas


def _ler_resumos(planilha: Any) -> list[ResumoUrna]:
    cols, linhas = _tabela(
        planilha,
        ("zona", "seção", "cargo", "aptos", "comparecimento", "válidos"),
    )
    c_zona = _indice(cols, "zona")
    c_secao = _indice(cols, "seção", "secao")
    c_codigo = _indice(cols, "cód. local", "código", "cod local")
    c_local = _indice(cols, "local de votação")
    c_agrega = _indice(cols, "agrega")
    c_cargo = _indice(cols, "cargo")
    c_aptos = _indice(cols, "aptos")
    c_comp = _indice(cols, "comparecimento")
    c_abst = _indice(cols, "abstenções", "abstencoes")
    c_nom = _indice(cols, "nominais")
    c_leg = _indice(cols, "legenda")
    c_val = _indice(cols, "válidos", "validos")
    c_bran = _indice(cols, "brancos")
    c_nul = _indice(cols, "nulos")
    c_total = _indice(cols, "total de votos")

    resumos: list[ResumoUrna] = []
    for linha in linhas:
        zona = _texto(_valor(linha, c_zona))
        secao = _texto(_valor(linha, c_secao))
        cargo = _texto(_valor(linha, c_cargo))
        if not zona.isdigit() or not secao.isdigit() or not cargo:
            continue
        resumos.append(
            ResumoUrna(
                zona=zona,
                secao=secao,
                codigo=_texto(_valor(linha, c_codigo)),
                local=_texto(_valor(linha, c_local)),
                agrega=_texto(_valor(linha, c_agrega)),
                cargo=cargo,
                aptos=_inteiro(_valor(linha, c_aptos)),
                comparecimento=_inteiro(_valor(linha, c_comp)),
                abstencoes=_inteiro(_valor(linha, c_abst)),
                nominais=_inteiro(_valor(linha, c_nom)),
                legenda=_inteiro(_valor(linha, c_leg)),
                validos=_inteiro(_valor(linha, c_val)),
                brancos=_inteiro(_valor(linha, c_bran)),
                nulos=_inteiro(_valor(linha, c_nul)),
                total=_inteiro(_valor(linha, c_total)),
            )
        )
    return resumos


def _ler_votos(planilha: Any) -> list[Voto]:
    cols, linhas = _tabela(
        planilha,
        ("zona", "seção", "cargo", "tipo", "votável", "votos"),
    )
    c_zona = _indice(cols, "zona")
    c_secao = _indice(cols, "seção", "secao")
    c_codigo = _indice(cols, "cód. local", "código", "cod local")
    c_local = _indice(cols, "local de votação")
    c_agrega = _indice(cols, "agrega")
    c_cargo = _indice(cols, "cargo")
    c_tipo = _indice(cols, "tipo")
    c_num = _indice(cols, "nº", "número", "numero")
    c_votavel = _indice(cols, "votável", "votavel")
    c_partido = _indice(cols, "partido")
    c_votos = _indice(cols, "votos")
    c_aptos = _indice(cols, "aptos")
    c_comp = _indice(cols, "comparecimento")
    c_abst = _indice(cols, "abstenções", "abstencoes")

    votos: list[Voto] = []
    for linha in linhas:
        zona = _texto(_valor(linha, c_zona))
        secao = _texto(_valor(linha, c_secao))
        cargo = _texto(_valor(linha, c_cargo))
        tipo = _texto(_valor(linha, c_tipo))
        if not zona.isdigit() or not secao.isdigit() or not cargo or not tipo:
            continue
        votos.append(
            Voto(
                zona=zona,
                secao=secao,
                codigo=_texto(_valor(linha, c_codigo)),
                local=_texto(_valor(linha, c_local)),
                agrega=_texto(_valor(linha, c_agrega)),
                cargo=cargo,
                tipo=tipo,
                numero=_texto(_valor(linha, c_num)),
                votavel=_texto(_valor(linha, c_votavel)),
                partido=_texto(_valor(linha, c_partido)),
                votos=_inteiro(_valor(linha, c_votos)),
                aptos=_inteiro(_valor(linha, c_aptos)),
                comparecimento=_inteiro(_valor(linha, c_comp)),
                abstencoes=_inteiro(_valor(linha, c_abst)),
            )
        )
    return votos


def ler_planilha(caminho: str | Path) -> PlanilhaBU:
    """Lê a planilha completa (metadados + locais + urnas + resumos + votos)."""

    alvo = Path(caminho)
    if not alvo.exists():
        raise PlanilhaError(f"Planilha não encontrada: {alvo}")
    wb = _abrir(alvo)
    try:
        meta = _ler_meta(_aba(wb, "Leia-me", "Leiame"))
        locais = _ler_locais(_aba(wb, "Locais de votação", "Locais de votacao"))
        urnas = _ler_urnas(_aba(wb, "Índice das urnas", "Indice das urnas"))
        resumos = _ler_resumos(_aba(wb, "Resumo por urna"))
        votos = _ler_votos(_aba(wb, "Dados completos"))
    finally:
        wb.close()

    if not locais:
        raise PlanilhaError("Nenhum local de votação encontrado na planilha.")
    if not votos:
        raise PlanilhaError("Nenhum voto encontrado na aba 'Dados completos'.")
    return PlanilhaBU(
        caminho=alvo, meta=meta, locais=locais, urnas=urnas, resumos=resumos, votos=votos
    )
